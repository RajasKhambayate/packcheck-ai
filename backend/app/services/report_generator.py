import os
from datetime import datetime

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                 TableStyle, Image as RLImage)
from docx import Document
from docx.shared import Pt, RGBColor, Inches

from app import models
from app.config import settings

STATUS_COLORS = {
    "compliant": colors.HexColor("#1B7A43"),
    "non_compliant": colors.HexColor("#C0392B"),
    "needs_review": colors.HexColor("#B8860B"),
}


def _ensure_dir():
    os.makedirs(settings.REPORT_DIR, exist_ok=True)


def generate_pdf_report(scan: models.ScanRecord) -> str:
    _ensure_dir()
    path = os.path.join(settings.REPORT_DIR, f"{scan.id}.pdf")

    doc = SimpleDocTemplate(path, pagesize=A4, topMargin=20 * mm, bottomMargin=20 * mm)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("TitleX", parent=styles["Title"], textColor=colors.HexColor("#1F2A44"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], textColor=colors.HexColor("#1F2A44"))
    normal = styles["Normal"]

    story = [
        Paragraph("PackCheck AI — Legal Metrology Compliance Report", title_style),
        Paragraph("Legal Metrology (Packaged Commodities) Rules, 2011", normal),
        Spacer(1, 10),
        Paragraph(f"Report generated: {datetime.utcnow().strftime('%d %b %Y, %H:%M UTC')}", normal),
        Spacer(1, 16),
        Paragraph("Product Details", h2),
    ]

    details_data = [
        ["Product Name", scan.product_name],
        ["Category", scan.category],
        ["Brand", scan.brand or "—"],
        ["Source", scan.source],
        ["Scan ID", scan.id],
        ["Scanned On", scan.created_at.strftime("%d %b %Y, %H:%M UTC")],
    ]
    t = Table(details_data, colWidths=[130, 330])
    t.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#555555")),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
    ]))
    story.append(t)
    story.append(Spacer(1, 16))

    story.append(Paragraph("Extracted Declarations", h2))
    decl_data = [
        ["Manufacturer / Packer / Importer", scan.manufacturer_name or "NOT DETECTED"],
        ["Manufacturer Address", scan.manufacturer_address or "NOT DETECTED"],
        ["Net Quantity", scan.net_quantity or "NOT DETECTED"],
        ["MRP", scan.mrp or "NOT DETECTED"],
        ["Mfg / Packing / Import Date", scan.mfg_date or "NOT DETECTED"],
        ["Consumer Care Details", scan.consumer_care or "NOT DETECTED"],
        ["Country of Origin", scan.country_of_origin or "NOT DETECTED / N.A."],
    ]
    t2 = Table(decl_data, colWidths=[180, 280])
    t2.setStyle(TableStyle([
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F2F4F8")),
    ]))
    story.append(t2)
    story.append(Spacer(1, 16))

    status_color = STATUS_COLORS.get(scan.status.value, colors.black)
    story.append(Paragraph("Compliance Verdict", h2))
    story.append(Paragraph(
        f'<font color="{status_color.hexval()}"><b>{scan.status.value.replace("_", " ").upper()}</b></font>'
        f'  —  Compliance Score: {scan.compliance_score}/100', normal))
    story.append(Spacer(1, 10))

    if scan.violations:
        story.append(Paragraph("Violations / Flags", h2))
        v_data = [["Rule", "Description", "Severity"]]
        for v in scan.violations:
            v_data.append([v.rule_code, v.rule_description, v.severity.value.upper()])
        t3 = Table(v_data, colWidths=[70, 320, 70])
        t3.setStyle(TableStyle([
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#DDDDDD")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F2A44")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ]))
        story.append(t3)
    else:
        story.append(Paragraph("No violations detected.", normal))

    story.append(Spacer(1, 20))
    story.append(Paragraph(
        "This report was generated automatically by PackCheck AI and is intended "
        "as a first-pass screening aid for enforcement officials. Findings marked "
        "NOT DETECTED indicate the automated system could not confidently read the "
        "declaration and require manual verification before any enforcement action.",
        ParagraphStyle("Disclaimer", parent=normal, fontSize=7.5, textColor=colors.grey)
    ))

    doc.build(story)
    return path


def generate_docx_report(scan: models.ScanRecord) -> str:
    _ensure_dir()
    path = os.path.join(settings.REPORT_DIR, f"{scan.id}.docx")

    document = Document()
    title = document.add_heading("PackCheck AI — Legal Metrology Compliance Report", level=1)
    title.runs[0].font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)

    document.add_paragraph("Legal Metrology (Packaged Commodities) Rules, 2011")
    document.add_paragraph(f"Report generated: {datetime.utcnow().strftime('%d %b %Y, %H:%M UTC')}")

    document.add_heading("Product Details", level=2)
    table = document.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in [
        ("Product Name", scan.product_name),
        ("Category", scan.category),
        ("Brand", scan.brand or "—"),
        ("Source", scan.source),
        ("Scan ID", scan.id),
        ("Scanned On", scan.created_at.strftime("%d %b %Y, %H:%M UTC")),
    ]:
        row = table.add_row().cells
        row[0].text = label
        row[1].text = str(value)

    document.add_heading("Extracted Declarations", level=2)
    table2 = document.add_table(rows=0, cols=2)
    table2.style = "Light Grid Accent 1"
    for label, value in [
        ("Manufacturer / Packer / Importer", scan.manufacturer_name or "NOT DETECTED"),
        ("Manufacturer Address", scan.manufacturer_address or "NOT DETECTED"),
        ("Net Quantity", scan.net_quantity or "NOT DETECTED"),
        ("MRP", scan.mrp or "NOT DETECTED"),
        ("Mfg / Packing / Import Date", scan.mfg_date or "NOT DETECTED"),
        ("Consumer Care Details", scan.consumer_care or "NOT DETECTED"),
        ("Country of Origin", scan.country_of_origin or "NOT DETECTED / N.A."),
    ]:
        row = table2.add_row().cells
        row[0].text = label
        row[1].text = str(value)

    document.add_heading("Compliance Verdict", level=2)
    p = document.add_paragraph()
    run = p.add_run(f"{scan.status.value.replace('_', ' ').upper()}  —  Score: {scan.compliance_score}/100")
    run.bold = True
    run.font.size = Pt(13)

    if scan.violations:
        document.add_heading("Violations / Flags", level=2)
        table3 = document.add_table(rows=1, cols=3)
        table3.style = "Light List Accent 2"
        hdr = table3.rows[0].cells
        hdr[0].text, hdr[1].text, hdr[2].text = "Rule", "Description", "Severity"
        for v in scan.violations:
            row = table3.add_row().cells
            row[0].text = v.rule_code
            row[1].text = v.rule_description
            row[2].text = v.severity.value.upper()
    else:
        document.add_paragraph("No violations detected.")

    document.add_paragraph()
    note = document.add_paragraph(
        "This report was generated automatically by PackCheck AI and is intended as a "
        "first-pass screening aid for enforcement officials. Findings require manual "
        "verification before any enforcement action."
    )
    note.runs[0].font.size = Pt(8)
    note.runs[0].font.color.rgb = RGBColor(0x80, 0x80, 0x80)

    document.save(path)
    return path
